"""
Excel de la consulta cruzada de stock: una fila por variante, una columna
por punto de venta — mismo contenido que la tabla de /consulta-stock, pero
sin paginar (`consulta_cruzada(tamano=None)`), porque acá se exporta todo lo
que matchea los filtros, no solo la página que se ve en pantalla.
"""

from io import BytesIO

from app.models.punto_de_venta import PuntoDeVenta


def generar_xls_consulta_stock(filas: list[dict], columnas: list[PuntoDeVenta]) -> bytes:
    """*filas* y *columnas* son el resultado de `servicio.consulta_cruzada`."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Stock por Local"

    negrita = Font(bold=True)

    # --- Cabecera de tabla ---
    encabezado = ["Código", "Descripción"] + [c.nombre for c in columnas]
    hoja.append(encabezado)
    for cell in hoja[hoja.max_row]:
        cell.font = negrita

    # --- Filas ---
    for fila in filas:
        codigo = fila["codigo_completo"] + (fila["verificador"] or "")
        desc = fila["descripcion"]
        if fila["descripcion_sufijo"]:
            desc += f" · {fila['descripcion_sufijo']}"
        stocks = fila["stocks"]
        hoja.append([codigo, desc] + [stocks.get(c.id, 0) for c in columnas])

    # --- Anchos de columna ---
    hoja.column_dimensions["A"].width = 16
    hoja.column_dimensions["B"].width = 40
    for i in range(len(columnas)):
        hoja.column_dimensions[get_column_letter(3 + i)].width = 12

    # Alinear las columnas de stock (numéricas) a la derecha.
    for row in hoja.iter_rows(min_row=1, max_row=hoja.max_row, min_col=3, max_col=2 + len(columnas)):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")

    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
