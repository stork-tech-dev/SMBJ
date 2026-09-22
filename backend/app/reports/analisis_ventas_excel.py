"""
Excel del reporte "Análisis de ventas por producto": una fila por
variante, sin paginar — se exporta todo lo que matchea los filtros, no
solo la página que se ve en pantalla.
"""

from io import BytesIO


def generar_xls_analisis_ventas(filas: list[dict]) -> bytes:
    """*filas* es el resultado de `servicio.analisis_por_producto`."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Análisis de ventas"

    negrita = Font(bold=True)

    encabezado = [
        "Código", "Descripción", "Categoría", "Proveedor",
        "Cant. 30 días", "Cant. 90 días", "Primera venta", "Última venta",
    ]
    hoja.append(encabezado)
    for cell in hoja[hoja.max_row]:
        cell.font = negrita

    for fila in filas:
        descripcion = fila["descripcion"]
        if fila["descripcion_sufijo"]:
            descripcion += f" · {fila['descripcion_sufijo']}"
        hoja.append([
            fila["codigo_completo"] + (fila["verificador"] or ""),
            descripcion,
            fila["categoria_nombre"],
            fila["proveedor_nombre"],
            fila["cantidad_30_dias"],
            fila["cantidad_90_dias"],
            fila["fecha_primera_venta"].strftime("%d/%m/%Y"),
            fila["fecha_ultima_venta"].strftime("%d/%m/%Y"),
        ])

    anchos = {
        "A": 16, "B": 40, "C": 20, "D": 20,
        "E": 14, "F": 14, "G": 14, "H": 14,
    }
    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho

    # Alinear las columnas numéricas (Cant. 30/90 días) a la derecha.
    for row in hoja.iter_rows(min_row=1, max_row=hoja.max_row, min_col=5, max_col=6):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")

    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
