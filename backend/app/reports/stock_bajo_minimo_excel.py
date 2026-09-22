"""
Excel del reporte "Productos bajo stock mínimo": una fila por combinación
de variante y ubicación, sin paginar — se exporta todo lo que matchea los
filtros, no solo la página que se ve en pantalla.
"""

from io import BytesIO

from app.models.stock import Stock


def generar_xls_stock_bajo_minimo(filas: list[Stock]) -> bytes:
    """*filas* es el resultado de `servicio.listar_stock(solo_bajo_minimo=True, ...)`."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Bajo stock mínimo"

    negrita = Font(bold=True)

    encabezado = [
        "Ubicación", "Código", "Descripción", "Categoría", "Proveedor",
        "Stock actual", "Mínimo",
    ]
    hoja.append(encabezado)
    for cell in hoja[hoja.max_row]:
        cell.font = negrita

    for fila in filas:
        producto = fila.variante.producto
        descripcion = producto.descripcion
        if fila.variante.descripcion_sufijo:
            descripcion += f" · {fila.variante.descripcion_sufijo}"
        hoja.append([
            fila.punto_de_venta.nombre,
            fila.variante.codigo_completo + (fila.variante.verificador or ""),
            descripcion,
            producto.categoria.nombre,
            producto.proveedor.nombre,
            fila.cantidad,
            fila.stock_minimo,
        ])

    anchos = {"A": 18, "B": 16, "C": 40, "D": 20, "E": 20, "F": 12, "G": 10}
    for columna, ancho in anchos.items():
        hoja.column_dimensions[columna].width = ancho

    # Alinear las columnas numéricas (Stock actual, Mínimo) a la derecha.
    for row in hoja.iter_rows(min_row=1, max_row=hoja.max_row, min_col=6, max_col=7):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")

    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
