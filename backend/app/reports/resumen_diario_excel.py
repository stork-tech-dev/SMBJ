"""
Excel del reporte "Resumen diario consolidado": una fila por turno, con su
local, franja horaria, cantidad de ventas y total. Sin paginar.
"""

from io import BytesIO


def generar_xls_resumen_diario(filas: list[dict]) -> bytes:
    """*filas* es el resultado de `servicio.resumen_diario_consolidado`."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Resumen diario"

    negrita = Font(bold=True)

    hoja.append(["Local", "Turno", "Cant. ventas", "Total"])
    for cell in hoja[hoja.max_row]:
        cell.font = negrita

    for fila in filas:
        desde = fila["fecha_apertura"].strftime("%H:%M")
        hasta = fila["fecha_cierre"].strftime("%H:%M") if fila["fecha_cierre"] else "en curso"
        hoja.append([
            fila["punto_de_venta_nombre"],
            f"{desde} – {hasta}",
            fila["cantidad_ventas"],
            float(fila["total"]),
        ])

    hoja.column_dimensions["A"].width = 28
    hoja.column_dimensions["B"].width = 20
    hoja.column_dimensions["C"].width = 14
    hoja.column_dimensions["D"].width = 16

    for row in hoja.iter_rows(min_row=2, max_row=hoja.max_row, min_col=3, max_col=4):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")

    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
