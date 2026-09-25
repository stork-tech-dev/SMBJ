"""
Excel de una tabla simple: encabezado en negrita, montos alineados a la
derecha, filas de subtotal en negrita y filas destacadas en rojo.

Lo usan los Reportes de Caja, que son todos tablas con la misma forma: cada
reporte arma sus filas y esto las escribe (Principio 2).
"""

from collections.abc import Iterable, Sequence
from decimal import Decimal
from io import BytesIO

ROJO = "F60509"  # --color-danger del design system


def generar_xls_tabla(
    titulo_hoja: str,
    encabezados: Sequence[str],
    filas: Iterable[Sequence],
    *,
    columnas_monto: Iterable[int] = (),
    filas_negrita: Iterable[int] = (),
    filas_destacadas: Iterable[int] = (),
    anchos: Sequence[int] | None = None,
) -> bytes:
    """
    *filas* son listas de valores ya en su tipo (Decimal se escribe como
    número). Los índices de `filas_negrita` / `filas_destacadas` son
    posiciones dentro de *filas* (0 = primera fila de datos, sin contar el
    encabezado); los de `columnas_monto`, posiciones de columna.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    hoja = wb.active
    hoja.title = titulo_hoja[:31]  # límite de Excel para el nombre de hoja

    hoja.append(list(encabezados))
    for celda in hoja[1]:
        celda.font = Font(bold=True)

    montos = set(columnas_monto)
    negritas = set(filas_negrita)
    destacadas = set(filas_destacadas)

    for i, fila in enumerate(filas):
        hoja.append([float(v) if isinstance(v, Decimal) else v for v in fila])
        fuente = Font(bold=i in negritas, color=ROJO if i in destacadas else None)
        for j, celda in enumerate(hoja[hoja.max_row]):
            celda.font = fuente
            if j in montos:
                celda.alignment = Alignment(horizontal="right")
                celda.number_format = "#,##0.00"

    for j, ancho in enumerate(anchos or [18] * len(encabezados)):
        hoja.column_dimensions[get_column_letter(j + 1)].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
