"""
Excel del reporte "Retiros de mercadería": lo que se descuenta del sueldo de
cada empleada. Primero las de esta empresa, después las de la otra (para
pasarle el dato). Sin paginar.
"""

from io import BytesIO


def generar_xls_retiros_mercaderia(retiros: list) -> bytes:
    """*retiros* es el resultado de `servicio.reporte(..., tamano=None)`."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Retiros de mercadería"
    negrita = Font(bold=True)

    hoja.append([
        "Empresa", "Empleada", "DNI", "Fecha", "Local", "Código", "Producto",
        "Precio lista", "Descuento %", "A descontar del sueldo",
    ])
    for cell in hoja[hoja.max_row]:
        cell.font = negrita

    for r in retiros:
        variante = r.variante
        producto = variante.producto.descripcion
        if variante.descripcion_sufijo:
            producto = f"{producto} — {variante.descripcion_sufijo}"
        hoja.append([
            "Esta empresa" if r.es_empresa_propia else "Otra empresa",
            r.nombre_empleada,
            r.empleada_dni or "",
            r.timestamp.strftime("%d/%m/%Y %H:%M"),
            r.punto_de_venta.nombre,
            variante.codigo_con_verificador,
            producto,
            float(r.precio_lista),
            float(r.descuento_aplicado),
            float(r.precio_con_descuento),
        ])

    for columna, ancho in zip("ABCDEFGHIJ", (14, 28, 12, 17, 22, 12, 40, 14, 12, 22)):
        hoja.column_dimensions[columna].width = ancho
    for row in hoja.iter_rows(min_row=2, max_row=hoja.max_row, min_col=8, max_col=10):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")

    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
