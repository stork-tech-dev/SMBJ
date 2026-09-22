"""
Excel del reporte "Resumen por punto de venta": una fila por local, con el
total y el porcentaje cobrado con cada medio de pago. Sin paginar.
"""

from io import BytesIO


def generar_xls_resumen_por_punto(filas: list[dict], columnas_medios: list[str]) -> bytes:
    """
    *filas* es el resultado de `servicio.resumen_por_punto_de_venta`.

    Las columnas de medio de pago son dinámicas: solo entran los medios que
    de verdad se usaron en el período filtrado.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Resumen por punto de venta"

    negrita = Font(bold=True)

    encabezado = ["Local", "Total"] + [f"% {medio}" for medio in columnas_medios]
    hoja.append(encabezado)
    for cell in hoja[hoja.max_row]:
        cell.font = negrita

    for fila in filas:
        hoja.append([
            fila["punto_de_venta_nombre"],
            float(fila["total"]),
            *[float(fila["porcentajes"].get(medio, 0)) for medio in columnas_medios],
        ])

    hoja.column_dimensions["A"].width = 28
    hoja.column_dimensions["B"].width = 16
    for i in range(len(columnas_medios)):
        columna = chr(ord("C") + i)
        hoja.column_dimensions[columna].width = 16

    ultima_columna = 2 + len(columnas_medios)
    for row in hoja.iter_rows(min_row=2, max_row=hoja.max_row, min_col=2, max_col=ultima_columna):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")
            if cell.column >= 3:
                cell.number_format = "0.0"

    buffer = BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()
